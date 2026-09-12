import io
from datetime import date

from docx import Document
from kall.models import CandidateProfile, Education, Employment, ResumeDocument, Skill, User
from kall.services.documents import _normalize_zip
from kall.services.resume_assembly import assemble_resume, layout_text
from kall.services.resume_render import (
    TREATMENTS,
    ordered_sections,
    render_docx,
    render_pdf,
    treatment_for,
)
from pypdf import PdfReader
from sqlmodel import Session, SQLModel, create_engine

TAILORED = [
    {"section": "summary", "text": "Quality leader with 15 years of experience.\n\nRole focus: Director at North, emphasizing test automation."},
    {"section": "achievement", "text": "Cut release defects 40% by rebuilding the regression suite."},
]


def _session() -> Session:
    engine = create_engine("sqlite://")
    SQLModel.metadata.create_all(engine)
    return Session(engine)


def _user_with_record(session: Session) -> User:
    user = User(email="ava@example.com", full_name="Ava Quality")
    session.add(user)
    session.commit()
    session.refresh(user)
    session.add(CandidateProfile(user_id=user.id, city="Vancouver", state_region="WA", linkedin_url="https://www.linkedin.com/in/ava"))
    session.add(Employment(user_id=user.id, employer="North", job_title="Director of Quality Engineering", location="Remote", start_date=date(2019, 3, 1), is_current=True, description="Led a 30-person QA org.\nBuilt CI gates for 40 services."))
    session.add(Employment(user_id=user.id, employer="Acme", job_title="QA Manager", start_date=date(2014, 1, 1), end_date=date(2019, 2, 1), description="Ran regression for the payments platform."))
    session.add(Education(user_id=user.id, institution="State University", degree="B.S.", major="Computer Science", graduation_date=date(2013, 6, 1)))
    session.add(Skill(user_id=user.id, name="Playwright", category="Automation", is_primary=True))
    session.add(Skill(user_id=user.id, name="Leadership", category="Management"))
    session.commit()
    return user


def test_layout_assembles_header_summary_jobs_skills_and_education_from_the_record() -> None:
    with _session() as session:
        user = _user_with_record(session)
        layout = assemble_resume(session, user.id, TAILORED)
    assert layout["name"] == "Ava Quality"
    assert layout["contact"] == ["ava@example.com", "Vancouver, WA", "linkedin.com/in/ava"]
    keys = [section["key"] for section in layout["sections"]]
    assert keys == ["summary", "experience", "achievements", "skills", "education"]
    summary = layout["sections"][0]
    assert summary["paragraphs"] == ["Quality leader with 15 years of experience.", "Role focus: Director at North, emphasizing test automation."]
    experience = layout["sections"][1]["entries"]
    assert experience[0]["title"] == "Director of Quality Engineering"
    assert experience[0]["dates"] == "Mar 2019 – Present"
    assert experience[0]["bullets"] == ["Led a 30-person QA org.", "Built CI gates for 40 services."]
    assert experience[1]["dates"] == "Jan 2014 – Feb 2019"
    assert layout["sections"][3]["groups"][0] == {"label": "Automation", "items": ["Playwright"]}
    assert layout["sections"][4]["entries"][0]["title"] == "B.S. in Computer Science"
    assert "ACHIEVEMENTS" in layout_text(layout).upper()


def test_layout_falls_back_to_the_parsed_resume_when_the_record_is_empty() -> None:
    with _session() as session:
        user = User(email="new@example.com", full_name="New Person")
        session.add(user)
        session.commit()
        session.refresh(user)
        resume = ResumeDocument(
            user_id=user.id, name="r.pdf", file_path="x", mime_type="application/pdf", byte_size=1,
            extracted_text="Summary\nSeasoned tester.\nExperience\nQA Lead at Acme 2015 - 2020\nSkills\nSelenium, Python\n",
        )
        session.add(resume)
        session.commit()
        layout = assemble_resume(session, user.id, [], resume)
    keys = [section["key"] for section in layout["sections"]]
    assert "experience" in keys
    experience = next(section for section in layout["sections"] if section["key"] == "experience")
    assert "QA Lead at Acme 2015 - 2020" in experience["paragraphs"]


def test_layout_reflows_word_per_line_extraction_before_falling_back() -> None:
    """Regression test: a resume uploaded before extracted_text was reflowed
    at upload time still has the old pypdf word-per-line bytes on disk --
    nothing repairs that row in place. Parsing it unreflowed put a single
    word ("Led") in as the whole experience entry and dropped the rest of
    the sentence, exactly like the "Experience" section that rendered one
    word per line in production."""
    with _session() as session:
        user = User(email="stale@example.com", full_name="Stale Upload")
        session.add(user)
        session.commit()
        session.refresh(user)
        resume = ResumeDocument(
            user_id=user.id, name="r.pdf", file_path="x", mime_type="application/pdf", byte_size=1,
            extracted_text=(
                "Summary\n\nSeasoned\ntester\nwith\nyears\nof\nexperience.\n\n"
                "Experience\n\nLed\nthe\nplatform\nteam\nat\nAcme.\n"
            ),
        )
        session.add(resume)
        session.commit()
        layout = assemble_resume(session, user.id, [], resume)
    experience = next(section for section in layout["sections"] if section["key"] == "experience")
    assert experience["paragraphs"] == ["Led the platform team at Acme."]


def test_every_template_renders_a_real_document_deterministically() -> None:
    with _session() as session:
        user = _user_with_record(session)
        layout = assemble_resume(session, user.id, TAILORED)
    for key in TREATMENTS:
        pdf = render_pdf(layout, key)
        assert pdf == render_pdf(layout, key), f"{key} pdf is not reproducible"
        text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(pdf)).pages)
        assert "Ava Quality" in text
        assert "Director of Quality Engineering" in text
        assert "Mar 2019" in text
        # ensure_artifact() normalizes the zip's entry timestamps after
        # rendering; that pass is what makes the bytes reproducible.
        docx = _normalize_zip(render_docx(layout, key))
        assert docx == _normalize_zip(render_docx(layout, key)), f"{key} docx is not reproducible"
        paragraphs = [p.text for p in Document(io.BytesIO(docx)).paragraphs]
        assert paragraphs[0] == "Ava Quality"
        assert any("Director of Quality Engineering" in p and "Mar 2019" in p for p in paragraphs)
        assert any(p == "Led a 30-person QA org." for p in paragraphs)


def test_templates_order_sections_differently() -> None:
    with _session() as session:
        user = _user_with_record(session)
        layout = assemble_resume(session, user.id, TAILORED)
    standard = [s["key"] for s in ordered_sections(layout, treatment_for("standard"))]
    executive = [s["key"] for s in ordered_sections(layout, treatment_for("executive"))]
    service = [s["key"] for s in ordered_sections(layout, treatment_for("service"))]
    assert standard.index("experience") < standard.index("achievements")
    assert executive.index("achievements") < executive.index("experience")
    assert service.index("skills") < service.index("experience")
